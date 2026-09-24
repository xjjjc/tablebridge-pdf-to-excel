"""Strict fixed-layout PDF extraction; no network calls or runtime AI."""
from __future__ import annotations

import io
import math
import re
from decimal import Decimal

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_PAGES = 50
MAX_ROWS = 10000


class ConversionError(ValueError):
    pass


def clean(value):
    return " ".join(str(value or "").split())


def validate_template(template):
    if not isinstance(template, dict):
        raise ConversionError("Template must be a JSON object. / 模板必须是 JSON 对象。")
    columns = template.get("columns")
    if not isinstance(columns, list) or not 1 <= len(columns) <= 30:
        raise ConversionError("Template needs 1–30 columns. / 模板需要 1–30 列。")
    names = []
    for column in columns:
        if not isinstance(column, dict) or not isinstance(column.get("name"), str):
            raise ConversionError("Each column needs a name. / 每列必须填写列名。")
        name = clean(column["name"])
        if not name or len(name) > 100:
            raise ConversionError("Column names must contain 1–100 characters.")
        if column.get("type", "text") not in ("text", "number"):
            raise ConversionError("Column type must be text or number.")
        if not isinstance(column.get("required", True), bool):
            raise ConversionError("required must be true or false.")
        names.append(name.casefold())
    if len(set(names)) != len(names):
        raise ConversionError("Column names must be unique. / 列名不能重复。")
    if set(names) & {"source file", "page", "table row"}:
        raise ConversionError("Source file, Page and Table row are reserved column names.")
    bbox = template.get("bbox", [0, 0, 1, 1])
    if not isinstance(bbox, list) or len(bbox) != 4 or any(
        isinstance(v, bool) or not isinstance(v, (float, int)) or not math.isfinite(v)
        for v in bbox
    ):
        raise ConversionError("bbox must be four finite numbers.")
    if not (0 <= bbox[0] < bbox[2] <= 1 and 0 <= bbox[1] < bbox[3] <= 1):
        raise ConversionError("bbox must be [left, top, right, bottom] fractions within 0–1.")
    size = template.get("page_size")
    if size is not None and (
        not isinstance(size, list) or len(size) != 2 or any(
            isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) or v <= 0
            for v in size
        )
    ):
        raise ConversionError("page_size must be [width, height] in PDF points.")
    checks = template.get("checks", [])
    if not isinstance(checks, list) or len(checks) > 10:
        raise ConversionError("checks must be a list of at most 10 multiplication checks.")
    number_names = {clean(c["name"]) for c in columns if c.get("type") == "number"}
    for check in checks:
        if not isinstance(check, dict) or set(check) != {"left", "right", "equals"}:
            raise ConversionError("Each check needs left, right and equals column names.")
        if any(not isinstance(v, str) or v not in number_names for v in check.values()):
            raise ConversionError("Arithmetic checks must reference numeric columns.")
    return template


def parse_number(text):
    # Explicit English-style decimal/grouping; reject ambiguous formats and units.
    if not re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?", text):
        raise ConversionError(f"Invalid number: {text!r}. Use 1234.56 or 1,234.56.")
    number = Decimal(text.replace(",", ""))
    if len(number.as_tuple().digits) > 15 or abs(number) > Decimal("1e15"):
        raise ConversionError("Number exceeds Excel's safe precision; use a text column.")
    return int(number) if number == number.to_integral_value() else float(number)


def extract_pdf(data: bytes, filename: str, template: dict):
    validate_template(template)
    if len(data) > MAX_FILE_BYTES:
        raise ConversionError("PDF exceeds 10 MB. / 单个 PDF 超过 10 MB。")
    if not data.startswith(b"%PDF-"):
        raise ConversionError("Not a valid PDF file. / 不是有效的 PDF 文件。")
    columns = template["columns"]
    headers = [clean(c["name"]) for c in columns]
    bbox = template.get("bbox", [0, 0, 1, 1])
    records, raw_records, warnings = [], [], []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if not 1 <= len(pdf.pages) <= MAX_PAGES:
                raise ConversionError("PDF must contain 1–50 pages. / 支持 1–50 页。")
            for page_number, page in enumerate(pdf.pages, 1):
                prefix = f"{filename}, page {page_number}"
                size = template.get("page_size")
                if size and (abs(page.width - size[0]) > 2 or abs(page.height - size[1]) > 2):
                    raise ConversionError(f"{prefix}: page size differs from template. / 页面尺寸不匹配。")
                region = page.crop((bbox[0] * page.width, bbox[1] * page.height,
                                    bbox[2] * page.width, bbox[3] * page.height))
                if not region.chars:
                    raise ConversionError(f"{prefix}: no selectable text. OCR/scans are not supported. / 无文字层，不支持扫描件。")
                tables = region.extract_tables({"vertical_strategy": "lines", "horizontal_strategy": "lines"})
                if len(tables) != 1:
                    raise ConversionError(f"{prefix}: expected exactly one ruled table in the selected area, found {len(tables)}. / 区域内必须恰好有一个带线表格。")
                table = tables[0]
                if len(table) < 2:
                    raise ConversionError(f"{prefix}: table has no data rows.")
                actual = [clean(v).casefold() for v in table[0]]
                if actual != [v.casefold() for v in headers]:
                    raise ConversionError(f"{prefix}: header mismatch. Expected {headers}; found {table[0]}. / 表头不匹配，请调整模板。")
                for row_number, row in enumerate(table[1:], 2):
                    if len(row) != len(columns):
                        raise ConversionError(f"{prefix}, row {row_number}: column count mismatch.")
                    if not any(clean(v) for v in row):
                        raise ConversionError(f"{prefix}, row {row_number}: empty row; check the PDF manually.")
                    record, raw = {}, {}
                    for column, value in zip(columns, row):
                        name = clean(column["name"])
                        text = clean(value)
                        if column.get("required", True) and not text:
                            raise ConversionError(f"{prefix}, row {row_number}: missing {name}.")
                        raw[name] = value or ""
                        try:
                            record[name] = parse_number(text) if text and column.get("type") == "number" else text
                        except ConversionError as exc:
                            raise ConversionError(f"{prefix}, row {row_number}, {name}: {exc}") from exc
                    for check in template.get("checks", []):
                        a, b, c = (record[check[k]] for k in ("left", "right", "equals"))
                        if any(v == "" for v in (a, b, c)):
                            warnings.append(f"{prefix}, row {row_number}: arithmetic check unavailable (blank value).")
                        elif abs(Decimal(str(a)) * Decimal(str(b)) - Decimal(str(c))) > Decimal("0.01"):
                            warnings.append(f"{prefix}, row {row_number}: {check['left']} × {check['right']} differs from {check['equals']}. Source values kept.")
                    for obj in (record, raw):
                        obj.update({"Source file": filename, "Page": page_number, "Table row": row_number})
                    records.append(record)
                    raw_records.append(raw)
                    if len(records) > MAX_ROWS:
                        raise ConversionError("Too many rows; split into smaller files.")
                page.close()
            return {"records": records, "raw_records": raw_records, "warnings": warnings,
                    "pages": len(pdf.pages), "filename": filename}
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError("PDF could not be read. Check for corruption or password protection. / 文件无法读取，请检查是否损坏或加密。") from exc


def export_excel(results, template):
    """Application exporter. PDF strings are always literal cells, never formulas."""
    validate_template(template)
    wb = Workbook()
    wb.remove(wb.active)
    headers = [clean(c["name"]) for c in template["columns"]] + ["Source file", "Page", "Table row"]
    number_names = {clean(c["name"]) for c in template["columns"] if c.get("type") == "number"}

    def populate(name, columns, rows):
        ws = wb.create_sheet(name)
        for i, values in enumerate([columns] + rows, 1):
            for j, value in enumerate(values, 1):
                if isinstance(value, str) and (len(value) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value)):
                    raise ConversionError("A cell contains unsupported characters or exceeds Excel's text limit.")
                cell = ws.cell(i, j, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="left" if isinstance(value, str) else "right")
                if i == 1:
                    cell.font = Font(name="Calibri", color="FFFFFF", bold=True)
                    cell.fill = PatternFill("solid", fgColor="163B40")
                else:
                    cell.font = Font(name="Calibri", size=11)
                    if i % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="EFF6F5")
                    if name == "Data" and columns[j - 1] in number_names and isinstance(value, (float, int)):
                        cell.number_format = '#,##0.00;[Red](#,##0.00);0.00'
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 28
        for j, column in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = 42 if column in ("Description", "Details") else 24 if column == "Source file" else 18
        for i, values in enumerate(rows, 2):
            line_count = max(sum(max(1, math.ceil(len(line) / max(8, int(ws.column_dimensions[get_column_letter(j)].width) - 2)))
                                 for line in str(value).split("\n"))
                             for j, value in enumerate(values, 1))
            ws.row_dimensions[i].height = max(23, 15 * line_count + 7)
        ws.sheet_view.showGridLines = False
        return ws

    populate("Data", headers, [[row.get(h, "") for h in headers] for r in results for row in r["records"]])
    populate("Raw text", headers, [[row.get(h, "") for h in headers] for r in results for row in r["raw_records"]])
    log_rows = [["Scope", "One ruled table per page; fixed header; no OCR. Review against the source before delivery."]]
    for result in results:
        log_rows.append([result["filename"], f"{result['pages']} pages; {len(result['records'])} rows; {len(result['warnings'])} arithmetic warnings"])
        log_rows.extend([["Check source", w] for w in result["warnings"]])
    populate("Review", ["Item", "Details"], log_rows)
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
