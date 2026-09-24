import copy
import io
import json
import unittest
from pathlib import Path

from openpyxl import load_workbook
from reportlab.pdfgen import canvas

from converter import ConversionError, extract_pdf, export_excel, validate_template
from scripts.create_samples import make_pdf

ROOT = Path(__file__).resolve().parents[1]


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.template = json.loads((ROOT / "templates/quotation.json").read_text(encoding="utf-8"))

    def extract(self, name):
        return extract_pdf((ROOT / "examples" / name).read_bytes(), name, self.template)

    def test_two_page_rows_types_and_provenance(self):
        result = self.extract("demo-quotation.pdf")
        self.assertEqual(result["pages"], 2)
        self.assertEqual(len(result["records"]), 8)
        self.assertEqual([r["Item"] for r in result["records"]], [f"{n:03}" for n in range(1, 9)])
        self.assertEqual(result["records"][4]["Page"], 2)
        self.assertEqual(result["records"][0]["Unit price"], 79.5)
        self.assertEqual(result["raw_records"][0]["Unit price"], "79.50")
        self.assertEqual(result["warnings"], [])

    def test_amount_warning_keeps_source_value(self):
        result = self.extract("demo-amount-warning.pdf")
        self.assertEqual(len(result["warnings"]), 1)
        self.assertEqual(result["records"][0]["Amount"], 160)

    def test_wrong_layout_rejected(self):
        with self.assertRaisesRegex(ConversionError, "header mismatch"):
            self.extract("demo-wrong-layout.pdf")

    def test_blank_page_rejected(self):
        buff = io.BytesIO();c = canvas.Canvas(buff);c.showPage();c.save()
        self.template.pop("page_size")
        with self.assertRaisesRegex(ConversionError, "no selectable text"):
            extract_pdf(buff.getvalue(), "blank.pdf", self.template)

    def test_invalid_input_rejected(self):
        with self.assertRaises(ConversionError):
            extract_pdf(b"not a PDF", "bad.pdf", self.template)

    def test_bad_templates_rejected(self):
        for field, value in [("columns", []), ("bbox", [0, .9, 1, .2]), ("page_size", [float("nan"), 792])]:
            with self.subTest(field=field):
                invalid = copy.deepcopy(self.template);invalid[field] = value
                with self.assertRaises(ConversionError):
                    validate_template(invalid)

    def test_workbook_preserves_ids_raw_text_and_numeric_types(self):
        result = self.extract("demo-quotation.pdf")
        wb = load_workbook(io.BytesIO(export_excel([result], self.template)))
        self.assertEqual(wb.sheetnames, ["Data", "Raw text", "Review"])
        self.assertEqual(wb["Data"].max_row, 9)
        self.assertEqual(wb["Data"]["A2"].value, "001")
        self.assertEqual(wb["Data"]["C2"].data_type, "n")
        self.assertEqual(wb["Raw text"]["D2"].value, "79.50")
        self.assertEqual(wb["Raw text"]["B7"].value, "Cable labels\nPack of 100")
        self.assertEqual(wb["Data"].freeze_panes, "A2")

    def test_formula_like_source_remains_literal(self):
        result = self.extract("demo-quotation.pdf")
        for rows in (result["records"], result["raw_records"]):
            rows[0]["Description"] = '=HYPERLINK("https://example.org", "text")'
        wb = load_workbook(io.BytesIO(export_excel([result], self.template)))
        self.assertEqual(wb["Data"]["B2"].data_type, "s")
        self.assertTrue(wb["Data"]["B2"].value.startswith("="))
        self.assertFalse(any(cell.data_type == "f" for sheet in wb for row in sheet for cell in row))

    def test_batch_keeps_both_sources(self):
        first = self.extract("demo-quotation.pdf")
        second = self.extract("demo-amount-warning.pdf")
        wb = load_workbook(io.BytesIO(export_excel([first, second], self.template)))
        self.assertEqual(wb["Data"].max_row, 11)
        self.assertEqual(wb["Data"]["F10"].value, "demo-amount-warning.pdf")


if __name__ == "__main__":
    unittest.main()
